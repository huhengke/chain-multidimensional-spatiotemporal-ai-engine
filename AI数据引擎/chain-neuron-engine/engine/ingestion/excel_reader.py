# -*- coding: utf-8 -*-
"""ExcelReader：openpyxl 流式只读；自动识别双表头（中文+英文码）。

另提供 `probe_file`：界面上传新数据源时的表头/列画像探测（注册向导用）。
"""
from __future__ import annotations

import itertools
import re
from collections.abc import Iterator
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

from engine.core.errors import ErrorCode, AppException
from engine.core.logging import get_logger
from engine.ingestion.cleaner import DAY_LEVEL_FORMATS

logger = get_logger("ingestion.reader")

# 探测参数
PROBE_SAMPLE_LIMIT = 200      # 画像统计的最大扫描行数
PROBE_PREVIEW_LIMIT = 20      # 预览返回行数
PROBE_COUNT_LIMIT = 500_000   # 行数统计上限（防超大文件拖慢探测）

_YYYYMM_RE = re.compile(r"^(\d{4})(\d{2})$")
_YYYYMMDD_RE = re.compile(r"^(\d{4})(\d{2})(\d{2})$")
# 纯日期/年月文本（允许日期单元格落到 Excel 后的 " 00:00:00" 时间尾巴）
_SEP_YM_RE = re.compile(r"^(\d{4})[-/.年](\d{1,2})月?(?:[ T]\d.*)?$")
_SEP_YMD_RE = re.compile(r"^(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?(?:[ T]\d.*)?$")
_DMY_RE = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})(?:[ T]\d.*)?$")
_SLUG_RE = re.compile(r"[^0-9a-zA-Z]+")
# 列名中的时间线索（中英文）
TIME_HINTS = ("月", "日期", "时间", "年月", "month", "date", "time", "period", "yyyymm")
CODE_HINTS = ("编码", "代码", "编号", "标识", "序号", "code", "_id", "id_", "id", "no", "key", "flag", "状态码")
METRIC_HINTS = ("金额", "数量", "业务数量", "业务", "收入", "价格", "单价", "总额", "余额", "amount", "qty", "quantity", "volume", "sales", "sale", "revenue", "price", "total", "value", "count")
# 可自动派生的时间层级（与 TimeDerivationEngine 派生名一致）
DERIVE_ORDER = ("year", "season", "month", "quarter", "half_year", "month_parity")
# 日粒度时间列还能继续派生周与日
DAY_DERIVE_ORDER = ("year", "season", "month", "quarter", "half_year", "month_parity", "week", "day")
DERIVE_LABELS = {
    "year": "年", "season": "季节", "month": "月", "quarter": "季度",
    "half_year": "上半年/下半年", "month_parity": "单月/双月", "solar_term": "二十四节气",
}
DERIVE_KEY_OF = {
    "year": "year", "season": "season", "month": "yyyymm", "quarter": "quarter",
    "half_year": "half_year", "month_parity": "month_parity", "solar_term": "solar_term",
}


def default_derive_for(time_format: str | None) -> list[str]:
    """声明的时间格式 → 建议派生层级（日粒度才能派生周/日）。"""
    if not time_format:
        return []
    fmt = str(time_format).strip().lower()
    return list(DAY_DERIVE_ORDER) if fmt in DAY_LEVEL_FORMATS else list(DERIVE_ORDER)


def slugify_name(text: str, fallback: str) -> str:
    """Excel 表头 → ASCII 列名建议（中文等无法转写时用 fallback）。"""
    slug = _SLUG_RE.sub("_", str(text or "")).strip("_").lower()
    if not slug or slug[0].isdigit():
        return fallback
    return slug


class _ColStat:
    """单列画像统计。"""

    def __init__(self) -> None:
        self.non_null = 0
        self.numeric = 0
        self.ym = 0         # 202401 形态
        self.ymd = 0        # 20240131 形态
        self.sep = 0        # 2024-01 / 2024年1月 形态
        self.sepday = 0     # 2024-01-31 / 2024-01-31 00:00:00 形态（日粒度）
        self.dmy = 0        # 2/1/2025 / 25/9/2025 15:24:43 形态（日在前）
        self.sample = None

    def feed(self, value) -> None:
        text = str(value).strip() if value is not None else ""
        if not text or text in ("*", "NULL", "null", "-", "None"):
            return
        self.non_null += 1
        if self.sample is None:
            self.sample = text
        if _to_decimal(text) is not None:
            self.numeric += 1
        m = _YYYYMM_RE.match(text)
        if m and 1 <= int(m.group(2)) <= 12:
            self.ym += 1
            return
        if _YYYYMMDD_RE.match(text):
            self.ymd += 1
            return
        if _SEP_YMD_RE.match(text):
            self.sepday += 1
            return
        if _SEP_YM_RE.match(text):
            self.sep += 1
            return
        m = _DMY_RE.match(text)
        if m and 1 <= int(m.group(2)) <= 12:
            self.dmy += 1

    def ratio(self, n: int) -> float:
        return round(n / self.non_null, 4) if self.non_null else 0.0

    def time_max(self) -> float:
        return max(self.ratio(self.ym), self.ratio(self.ymd),
                   self.ratio(self.sep), self.ratio(self.sepday),
                   self.ratio(self.dmy))


def _to_decimal(text: str) -> Decimal | None:
    try:
        d = Decimal(text.replace(",", ""))
    except (InvalidOperation, AttributeError):
        return None
    return d if d.is_finite() else None


def _preview_cell(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _guess_time_format(st: _ColStat) -> str:
    if st.ratio(st.ym) >= 0.9:
        return "yyyymm"
    if st.ratio(st.ymd) >= 0.9:
        return "yyyymmdd"
    if st.ratio(st.sepday) >= 0.9:
        return "yyyy-mm-dd"
    if st.ratio(st.sep) >= 0.9:
        return "yyyy-mm"
    if st.ratio(st.dmy) >= 0.9:
        return "d/m/yyyy"
    return "yyyymm"


def _is_identifier_hint(text: str) -> bool:
    text = str(text or "").strip().lower()
    if any(token in text for token in ("编码", "代码", "编号", "标识", "序号", "状态码")):
        return True
    if re.search(r"(?:^|[_\-\s])(id|code|no|key|flag)(?:$|[_\-\s])", text):
        return True
    return text.endswith(("id", "code", "_no", "no", "_key", "key", "flag"))


def _guess_role(st: _ColStat, hint_text: str) -> str:
    """建议字段角色：time / metric / dimension。"""
    if st.non_null == 0:
        return "ignore"
    lower_hint = hint_text.lower()
    hint = any(h in lower_hint for h in TIME_HINTS)
    if _is_identifier_hint(lower_hint):
        return "dimension"
    if st.time_max() >= 0.9:
        return "time"
    if hint and st.ratio(st.numeric) >= 0.9:
        return "time"
    if st.ratio(st.numeric) >= 0.9 and any(h in lower_hint for h in METRIC_HINTS):
        return "metric"
    return "dimension"


def probe_file(file_path: str | Path, sheet_name: str | None = None) -> dict:
    """探测 Excel 表头模式与列画像，给出建议字段角色（注册向导第一步）。

    复用 `ExcelReader` 的表头识别规则，保证"探测所见"与"导入所用"完全一致。
    """
    path = Path(file_path)
    ExcelReader(path, sheet_name)  # 复用存在性/后缀校验
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header_cn = next(rows, None)
        if header_cn is None:
            raise AppException(
                ErrorCode.IMPORT_FILE_ERROR[0], "Excel 文件为空", http_status=400
            )
        header_second = next(rows, None)
        if header_second is None:
            header_second = header_cn

        if ExcelReader._looks_like_codes(header_cn) or not ExcelReader._looks_like_codes(header_second):
            header_mode = "single"
            header = [ExcelReader._norm(c) for c in header_cn]
            cn_header = ["" for _ in header]
            first_data = header_second
        else:
            header_mode = "double"
            header = [ExcelReader._norm(c) for c in header_second]
            cn_header = [ExcelReader._norm(c) for c in header_cn]
            first_data = None
        header_rows = 1 if header_mode == "single" else 2

        stats = [_ColStat() for _ in header]
        preview: list[dict] = []
        row_count = 0
        capped = False
        # 行数优先取工作表维度（快速）；取不到时按实际扫描计数（上限 PROBE_COUNT_LIMIT）
        max_row = getattr(ws, "max_row", None)
        estimated = (
            int(max_row) - header_rows
            if isinstance(max_row, int) and max_row > header_rows
            else None
        )
        stream = rows if first_data is None else itertools.chain([first_data], rows)
        for raw in stream:
            if raw is None:
                continue
            if all(v is None or str(v).strip() == "" for v in raw):
                continue
            row_count += 1
            if row_count > PROBE_COUNT_LIMIT:
                capped = True
                row_count = PROBE_COUNT_LIMIT
                break
            if row_count <= PROBE_SAMPLE_LIMIT:
                for i in range(len(header)):
                    stats[i].feed(raw[i] if i < len(raw) else None)
            if len(preview) < PROBE_PREVIEW_LIMIT:
                preview.append({
                    header[i]: _preview_cell(raw[i] if i < len(raw) else None)
                    for i in range(len(header))
                })
            if estimated is not None and row_count >= PROBE_SAMPLE_LIMIT:
                break  # 画像与预览已足够，行数用工作表维度估计
        if estimated is not None and row_count >= PROBE_SAMPLE_LIMIT:
            row_count = max(estimated, row_count)
        else:
            estimated = None

        used: set[str] = set()
        columns: list[dict] = []
        time_candidates: list[tuple[float, dict]] = []  # (非空覆盖率, 列画像)
        sampled_total = max(min(row_count, PROBE_SAMPLE_LIMIT), 1)  # 画像实际扫描行数
        for i, source in enumerate(header):
            st = stats[i]
            role = _guess_role(st, f"{source}{cn_header[i]}")
            name = slugify_name(source or cn_header[i], f"field{i + 1}")
            base, n = name, 2
            while name in used:
                name = f"{base}_{n}"
                n += 1
            used.add(name)
            item = {
                "index": i,
                "source": source,
                "cn": cn_header[i],
                "sample": st.sample,
                "non_null": st.non_null,
                "numeric_ratio": st.ratio(st.numeric),
                "time_ratio": st.time_max(),
                "guess_role": role,
                "guess_name": name,
                "guess_label": cn_header[i] or source,
            }
            if role == "time":
                item["guess_time_format"] = _guess_time_format(st)
                time_candidates.append((min(st.non_null / sampled_total, 1.0), item))
            columns.append(item)

        # 时间列取非空覆盖最全者（覆盖率相差 5% 以内时保持列序优先），其余降级为维度
        time_col: dict | None = None
        if time_candidates:
            best = max(c[0] for c in time_candidates)
            threshold = max(0.0, best - 0.05)
            for coverage, item in time_candidates:
                if coverage >= threshold:
                    time_col = item
                    break
            for _, item in time_candidates:
                if item is not time_col:
                    item["guess_role"] = "dimension"

        default_derive = default_derive_for(time_col["guess_time_format"] if time_col else None)
        logger.info("Excel 探测：%s mode=%s rows=%s cols=%s",
                    path.name, header_mode, row_count, len(header))
        return {
            "file_name": path.name,
            "file_token": path.name,
            "sheet": ws.title,
            "sheets": list(wb.sheetnames),
            "header_mode": header_mode,
            "header_rows": header_rows,
            "row_count": row_count,
            "row_count_estimated": estimated is not None,
            "row_count_capped": capped,
            "columns": columns,
            "preview": preview,
            "time_column": time_col["source"] if time_col else None,
            "time_format": time_col["guess_time_format"] if time_col else "yyyymm",
            "default_derive": default_derive,
        }
    finally:
        wb.close()


class ExcelReader:
    def __init__(self, file_path: str | Path, sheet_name: str | None = None):
        self.path = Path(file_path)
        if not self.path.exists():
            raise AppException(
                ErrorCode.IMPORT_FILE_ERROR[0],
                "导入文件不存在",
                str(self.path),
                http_status=400,
            )
        if self.path.suffix.lower() != ".xlsx":
            raise AppException(
                ErrorCode.UNSUPPORTED_FILE[0],
                "仅支持 .xlsx 文件",
                self.path.suffix,
                http_status=400,
            )
        self.sheet_name = sheet_name

    def read(self) -> tuple[list[str], Iterator[tuple[int, dict]]]:
        """返回 (英文码列名列表, 迭代器(行号, 行字典))。

        双表头规则：第 1 行中文名（忽略），第 2 行英文码（作为键）；数据从第 3 行起。
        单表头规则：第 1 行即英文码，数据从第 2 行起（如维表/主数据文件）。
        """
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        ws = wb[self.sheet_name] if self.sheet_name else wb[wb.sheetnames[0]]

        rows = ws.iter_rows(values_only=True)
        try:
            header_cn = next(rows)
        except StopIteration:
            wb.close()
            raise AppException(
                ErrorCode.IMPORT_FILE_ERROR[0],
                "Excel 文件为空",
                http_status=400,
            )
        try:
            header_en = next(rows)
        except StopIteration:
            wb.close()
            header_en = header_cn

        # 双表头需满足：首行不是英文码表头，且第二行是英文码表头。
        # 首行即英文码 → 单表头（避免把首个数据行误当表头而丢一行）。
        if self._looks_like_codes(header_cn) or not self._looks_like_codes(header_en):
            # 单表头：数据从第 2 行开始
            columns = [self._norm(c) for c in header_cn]
            first_data_row = header_en
            start_idx = 2
        else:
            columns = [self._norm(c) for c in header_en]
            first_data_row = None
            start_idx = 3

        def _iter() -> Iterator[tuple[int, dict]]:
            row_idx = start_idx
            try:
                if first_data_row is not None:
                    yield row_idx, self._to_dict(columns, first_data_row)
                    row_idx += 1
                for raw in rows:
                    if raw is None:
                        continue
                    # 全空行跳过
                    if all(v is None or str(v).strip() == "" for v in raw):
                        row_idx += 1
                        continue
                    yield row_idx, self._to_dict(columns, raw)
                    row_idx += 1
            finally:
                wb.close()

        logger.info("Excel 打开：%s columns=%s", self.path.name, columns)
        return columns, _iter()

    @staticmethod
    def _norm(v) -> str:
        return str(v).strip() if v is not None else ""

    @staticmethod
    def _looks_like_codes(row: tuple) -> bool:
        """判断一行是否为英文码表头。

        要求非空值几乎全为 ASCII 标识符（字母/下划线开头，仅含字母数字下划线），
        以排除「以数值/中文/日期为主的」数据行（旧版仅看 ASCII 占比，会把首个
        数据行误判为表头）。
        """
        vals = [str(v).strip() for v in row if v is not None and str(v).strip() != ""]
        if not vals:
            return False

        def _is_code(v: str) -> bool:
            head = v[0]
            if not (head.isascii() and (head.isalpha() or head == "_")):
                return False
            return all(c.isascii() and (c.isalnum() or c == "_") for c in v)

        return sum(1 for v in vals if _is_code(v)) / len(vals) > 0.8

    @staticmethod
    def _to_dict(columns: list[str], raw: tuple) -> dict:
        return {columns[i]: raw[i] if i < len(raw) else None for i in range(len(columns))}
